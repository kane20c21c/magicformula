#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
yz_group.py — YZ 변동성 고·중·저 그룹 (연도 기준, 균등 3분할)
==============================================================
케인 지시 2026-09-10: "YZ 변동성의 26년도(8개월)을 기준으로 고중저 그룹을 나눌 수 있지?"

무엇을 만드나
  기준 연도의 데이터**만으로** 종목별 Yang-Zhang 변동성을 계산해 순위 균등 3분할.
  산출물 둘 —
    · JSON  `output/classification/yz_group_<연도>.json`  (코드가 읽는 정본)
    · Excel `docs/holdtiming/results/YZ그룹_<연도>.xlsx`   (사람이 보는 명단 + 진단)

왜 '그 해 데이터만' 인가
  저장된 `YZ_60` 은 60거래일 롤링이라 2026-01~03 값이 2025년 데이터를 물고 있다.
  "2026년 기준" 이라 부르려면 창이 연도 밖으로 새면 안 된다.
  (참고: 두 방식의 순위상관은 0.964 로 사실상 같다 — 그래도 이름과 내용을 맞춘다.)

⚠⚠ 읽는 사람이 반드시 알아야 할 것 (JSON·엑셀에도 박아 둔다)
  1. **이건 절대 수준이 아니라 그 해 안에서의 서열이다.** 시장 변동성은 통째로
     움직인다 — 2023-04 시장 중앙 0.0221 → 2026-08 0.0588 (2.66배). 그래서
     **2026년의 '저' 가 2023년의 '고' 보다 변동적일 수 있다.** 연도 간 비교 금지.
  2. **서열은 영구가 아니다.** 중첩 0 실측으로 순위상관이 3개월 0.796 → 12개월
     0.631 → 21개월 0.570 으로 감쇠한다. 그룹을 고정해 쓰면 시간이 갈수록 어긋난다
     — 아래 '안정성 진단' 이 그 크기를 숫자로 준다.
  3. BM(102110 TIGER 200)은 제외한다 — 지수 ETF 라 개별 종목 서열에 섞으면 안 된다.

실행
  python3 scripts/yz_group.py [--year 2026] [--cuts 0.333 0.667]
"""
from __future__ import annotations

import argparse
import json
import sys
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

HERE = Path(__file__).resolve()
REPO = HERE.parent.parent
STOLAB = REPO.parent
BM = "102110"
LABELS = ["저", "중", "고"]
MIN_DAYS_RATIO = 0.5          # 그 해 거래일의 이 비율 미만이면 판정 보류


def load_raw() -> pd.DataFrame:
    vault = STOLAB / "longlivevault" / "data" / "ohlcv"
    parts = [pd.read_parquet(vault / n) for n in ("core.parquet", "extend.parquet")
             if (vault / n).exists()]
    if not parts:
        raise SystemExit(f"❌ OHLCV parquet 없음: {vault}")
    df = pd.concat(parts, ignore_index=True)
    df["Date"] = pd.to_datetime(df["Date"])
    df["Ticker"] = df["Ticker"].astype(str).str.zfill(6)
    return df


def _piv(df: pd.DataFrame, col: str) -> pd.DataFrame:
    return (df.pivot_table(index="Date", columns="Ticker", values=col, aggfunc="last")
            .astype(float))


def yz_for(df: pd.DataFrame, start: str, end: str, drop_bm=True) -> tuple[pd.Series, pd.Series]:
    """구간 데이터**만으로** 종목별 YZ σ (일간) + 유효 거래일수."""
    O, H, L, C = (_piv(df, c) for c in ("Open", "High", "Low", "Close"))
    if drop_bm:
        O, H, L, C = (t.drop(columns=[BM], errors="ignore") for t in (O, H, L, C))
    # ⚠ o_t 는 전일 종가를 쓰므로 shift 는 구간을 자르기 **전에** 해야 한다.
    o_all, cc_all = np.log(O / C.shift(1)), np.log(C / O)
    u, dn = np.log(H / O), np.log(L / O)
    rs_all = u * (u - cc_all) + dn * (dn - cc_all)

    s = C.loc[start:end].index
    n = len(s)
    if n < 30:
        raise SystemExit(f"❌ 구간 거래일이 {n}일뿐 — 판정 불가")
    k = 0.34 / (1.34 + (n + 1) / (n - 1))
    o, cc, rs = o_all.loc[s], cc_all.loc[s], rs_all.loc[s]
    sig = np.sqrt(o.var(ddof=1) + k * cc.var(ddof=1) + (1 - k) * rs.mean())
    days = C.loc[s].notna().sum()
    return sig, days


def assign(sig: pd.Series, days: pd.Series, n_days: int,
           cuts=(1 / 3, 2 / 3)) -> pd.DataFrame:
    """균등 3분할. 표본이 얇은 종목은 판정 보류(None) — 순위 계산에서도 뺀다."""
    ok = sig.notna() & (days >= n_days * MIN_DAYS_RATIO)
    r = sig[ok].rank(pct=True)
    grp = pd.Series(np.select([r <= cuts[0], r <= cuts[1]], [LABELS[0], LABELS[1]],
                              default=LABELS[2]), index=r.index)
    out = pd.DataFrame({"YZ": sig, "관측일": days})
    out["순위"] = sig[ok].rank(ascending=False).astype("Int64")   # 1 = 가장 변동적
    out["백분위"] = (r * 100).round(1)
    out["그룹"] = grp
    out["연율%"] = (sig * np.sqrt(252) * 100).round(1)
    return out


def meta(df: pd.DataFrame, idx) -> pd.DataFrame:
    sys.path.insert(0, str(STOLAB / "longlivevault"))
    from stolab_data import TICKER_LIST, EXTEND_LIST, CORE_TICKERS  # noqa: E402
    m = {t[0]: (t[1], t[2]) for t in list(TICKER_LIST) + list(EXTEND_LIST)}
    nm = df.dropna(subset=["Name"]).sort_values("Date").groupby("Ticker")["Name"].last()
    mc = df.groupby("Ticker")["MarketCap"].median()
    return pd.DataFrame({
        "종목명": [m.get(t, ("", ""))[0] or nm.get(t, t) for t in idx],
        "섹터": [m.get(t, ("", "기타"))[1] for t in idx],
        "소속": ["core" if t in CORE_TICKERS else "extend" for t in idx],
        "시총(조)": [round(mc.get(t, np.nan) / 1e12, 2) if pd.notna(mc.get(t, np.nan))
                   else np.nan for t in idx],
    }, index=idx)


def stability(df: pd.DataFrame, year: int, cuts) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """이 그룹이 과거에도 같았나 — 연도별로 따로 나눠 보고 이탈률을 잰다.

    ⚠ 각 연도를 **그 해 데이터만으로** 다시 3분할한다. 절대 수준이 아니라 서열
      비교라야 의미가 있기 때문이다(연도 간 시장 변동성이 2.66배 차이난다).
    """
    years = [y for y in range(2023, year + 1)]
    g = {}
    for y in years:
        try:
            sig, days = yz_for(df, f"{y}-01-01", f"{y}-12-31")
            nd = len(_piv(df, "Close").loc[f"{y}-01-01":f"{y}-12-31"])
            g[y] = assign(sig, days, nd, cuts)["그룹"]
        except SystemExit:
            continue
    cur = g[year]
    rows = []
    for y in years:
        if y == year or y not in g:
            continue
        m = cur.notna() & g[y].notna()
        same = (cur[m] == g[y][m]).mean()
        # 정반대(고↔저)로 건너뛴 비율
        flip = (((cur[m] == "고") & (g[y][m] == "저")) |
                ((cur[m] == "저") & (g[y][m] == "고"))).mean()
        rows.append({"비교 연도": y, f"{year} 그룹과 일치": same,
                     "고↔저 정반대": flip, "대상 종목": int(m.sum())})
    keep = pd.DataFrame(rows)

    # 실전 관점 — 직전 연도로 그룹을 짰다면 올해 얼마나 맞았나
    prev = year - 1
    oos = {}
    if prev in g:
        m = cur.notna() & g[prev].notna()
        oos = {"기준": prev, "적중": float((cur[m] == g[prev][m]).mean()),
               "n": int(m.sum())}
        ct = pd.crosstab(g[prev][m], cur[m], normalize="index") * 100
        ct = ct.reindex(index=LABELS, columns=LABELS).fillna(0)
        ct.index = [f"{prev}년 {i}" for i in ct.index]
        ct.columns = [f"→{year}년 {c}" for c in ct.columns]
    else:
        ct = pd.DataFrame()
    return keep, ct.reset_index().rename(columns={"index": "이전"}), oos


def write_xlsx(out: Path, tbl: pd.DataFrame, keep: pd.DataFrame, trans: pd.DataFrame,
               oos: dict, year: int, info: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    hdr = PatternFill("solid", fgColor="FF37474F")
    thin = Side(style="thin", color="FFBDBDBD")
    med = Side(style="medium", color="FF37474F")
    FILL = {"고": "FFFFEBEE", "중": "FFF5F5F5", "저": "FFE3F2FD"}   # 고 연빨강 / 저 연파랑
    COLOR = {"고": "FFEF5350", "중": "FF616161", "저": "FF1976D2"}

    wb = Workbook()

    def sheet(name, widths):
        ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] else wb.active
        ws.title = name
        ws.sheet_view.showGridLines = False
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        return ws

    def note(ws, row, text, bold=False, size=10, color="FF000000"):
        c = ws.cell(row=row, column=1, value=text)
        c.font = Font(name=F, size=size, bold=bold, color=color)
        return row + 1

    def table(ws, df, r0, fmt, group_col=None):
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(row=r0, column=j, value=col)
            c.fill, c.font = hdr, Font(name=F, bold=True, color="FFFFFFFF", size=10)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.border = Border(bottom=med, left=thin, right=thin)
        ws.row_dimensions[r0].height = 28
        for i, (_, rec) in enumerate(df.iterrows()):
            g = rec.get(group_col) if group_col else None
            for j, col in enumerate(df.columns, 1):
                v = rec[col]
                c = ws.cell(row=r0 + 1 + i, column=j,
                            value=(None if (isinstance(v, float) and np.isnan(v)) else v))
                c.font = Font(name=F, size=10, bold=(col == group_col),
                              color=COLOR.get(g, "FF000000") if col == group_col else "FF000000")
                c.number_format = fmt.get(col, "General")
                c.alignment = Alignment(horizontal="left" if j == 1 else "right")
                c.border = Border(left=thin, right=thin, bottom=thin)
                if g in FILL:
                    c.fill = PatternFill("solid", fgColor=FILL[g])
        return r0 + len(df) + 1

    # ── 명단 ──
    ws = sheet(f"{year}_그룹", [8, 16, 14, 8, 6, 10, 9, 8, 9, 10, 9])
    r = note(ws, 1, f"YZ 변동성 고·중·저 그룹 — {year}년 기준", bold=True, size=14)
    r = note(ws, r, f"{info['start']} ~ {info['end']} · {info['n_days']}거래일 · "
                    f"{info['n']}종목 (BM 제외) · 균등 3분할 · "
                    f"생성기 scripts/yz_group.py", size=9, color="FF757575")
    r += 1
    for i, c in enumerate(CAVEATS, 1):
        r = note(ws, r, f"⚠ {i}. {c}", size=9, color="FF6D4C41")
    r += 1
    cols = ["티커", "종목명", "섹터", "소속", "그룹", "YZ", "연율%", "순위", "백분위",
            "시총(조)", "관측일"]
    r = table(ws, tbl[cols], r, {"YZ": "0.0000", "연율%": "0.0", "순위": "#,##0",
                                 "백분위": "0.0", "시총(조)": "#,##0.00",
                                 "관측일": "#,##0"}, group_col="그룹")
    ws.freeze_panes = ws.cell(row=r - len(tbl), column=3)

    # ── 요약·진단 ──
    ws = sheet("요약_안정성", [26, 14, 14, 14, 12])
    r = note(ws, 1, f"그룹 요약 및 안정성 진단 — {year}년 기준", bold=True, size=14)
    r += 1
    g = tbl.groupby("그룹")
    summ = pd.DataFrame({
        "그룹": LABELS,
        "종목수": [int((tbl["그룹"] == k).sum()) for k in LABELS],
        "YZ 중앙": [g["YZ"].median().get(k, np.nan) for k in LABELS],
        "연율% 중앙": [g["연율%"].median().get(k, np.nan) for k in LABELS],
        "시총(조) 중앙": [g["시총(조)"].median().get(k, np.nan) for k in LABELS],
    })
    r = table(ws, summ, r, {"종목수": "#,##0", "YZ 중앙": "0.0000",
                            "연율% 중앙": "0.0", "시총(조) 중앙": "#,##0.00"},
              group_col="그룹")
    r += 1

    if not keep.empty:
        r = note(ws, r, "이 그룹이 과거에도 같았나 — 각 연도를 그 해 데이터만으로 다시 3분할",
                 bold=True, size=12)
        r = note(ws, r, "   무작위면 일치 33% · 고↔저 정반대 22%. "
                        "연도가 멀수록 일치가 떨어지는 것이 정상이다.", size=9, color="FF757575")
        r = table(ws, keep, r + 1,
                  {f"{year} 그룹과 일치": "0.0%", "고↔저 정반대": "0.0%",
                   "비교 연도": "0", "대상 종목": "#,##0"})
        r += 1

    if oos:
        r = note(ws, r, f"★ 실전 관점 — {oos['기준']}년 기준으로 그룹을 짰다면 "
                        f"{year}년에 얼마나 맞았나", bold=True, size=12)
        r = note(ws, r, f"   적중 **{oos['적중']:.1%}** (n={oos['n']}) — 무작위 33.3%. "
                        "그룹을 1년 고정해 쓸 때의 신뢰도가 이 값이다.", size=10)
        if not trans.empty:
            r = table(ws, trans, r + 1, {c: "0.0" for c in trans.columns if c != "이전"})
            r = note(ws, r, "   행 = 이전 연도 그룹, 열 = 올해 어디로 갔나 (%)",
                     size=9, color="FF757575")
        r += 1
    r = note(ws, r, "⚠ 이 진단은 '그룹을 고정해 쓸 때 얼마나 어긋나는가' 만 말한다. "
                    "어느 그룹이 수익률에 유리한지는 여기서 판정하지 않는다.",
             size=9, color="FF6D4C41")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument("--cuts", type=float, nargs=2, default=[1 / 3, 2 / 3])
    a = ap.parse_args()

    print("· 패널 로드")
    df = load_raw()
    start, end = f"{a.year}-01-01", f"{a.year}-12-31"
    C = _piv(df, "Close").loc[start:end]
    n_days = len(C)
    print(f"· {a.year}년 구간 {C.index.min().date()} ~ {C.index.max().date()} ({n_days}거래일)")

    sig, days = yz_for(df, start, end)
    tbl = assign(sig, days, n_days, tuple(a.cuts)).join(meta(df, sig.index))
    tbl.index.name = "티커"
    tbl = tbl.reset_index().sort_values("순위")
    hold = tbl["그룹"].isna().sum()
    if hold:
        print(f"  ⚠ 표본 부족으로 판정 보류 {hold}종목")
    print(f"· 그룹: " + " · ".join(f"{k} {(tbl['그룹']==k).sum()}" for k in LABELS))

    print("· 안정성 진단")
    keep, trans, oos = stability(df, a.year, tuple(a.cuts))
    if oos:
        print(f"  {oos['기준']}년 기준으로 짰다면 {a.year}년 적중 {oos['적중']:.1%} "
              f"(무작위 33.3%)")

    js = REPO / "output" / "classification" / f"yz_group_{a.year}.json"
    js.parent.mkdir(parents=True, exist_ok=True)
    js.write_text(json.dumps({
        "generated_at": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M KST"),
        "generated_by": "MagicFormula/scripts/yz_group.py",
        "reproduce": f"python3 scripts/yz_group.py --year {a.year}",
        "method": "Yang-Zhang σ, 구간 데이터만으로 계산, 순위 균등 3분할",
        "window": {"start": str(C.index.min().date()), "end": str(C.index.max().date()),
                   "trading_days": n_days},
        "cuts": list(a.cuts), "labels": LABELS, "excluded": [BM],
        "caveats": CAVEATS,
        "stability": {"out_of_sample": oos,
                      "vs_prior_years": keep.to_dict("records") if not keep.empty else []},
        "groups": {k: sorted(tbl.loc[tbl["그룹"] == k, "티커"]) for k in LABELS},
        "detail": {r["티커"]: {"name": r["종목명"], "group": r["그룹"],
                              "yz": None if pd.isna(r["YZ"]) else round(r["YZ"], 6),
                              "rank": None if pd.isna(r["순위"]) else int(r["순위"]),
                              "sector": r["섹터"]}
                   for _, r in tbl.iterrows()},
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"✅ JSON  {js}")

    xl = REPO / "docs" / "holdtiming" / "results" / f"YZ그룹_{a.year}.xlsx"
    write_xlsx(xl, tbl, keep, trans, oos, a.year,
               {"start": str(C.index.min().date()), "end": str(C.index.max().date()),
                "n_days": n_days, "n": int(tbl["그룹"].notna().sum())})
    print(f"✅ Excel {xl}")
    return 0


CAVEATS = [
    "이건 절대 수준이 아니라 그 해 안에서의 서열이다. 시장 변동성은 통째로 움직인다"
    " — 2023-04 시장 중앙 0.0221 → 2026-08 0.0588 (2.66배).",
    "그래서 2026년의 '저' 가 2023년의 '고' 보다 변동적일 수 있다 — 연도 간 비교 금지.",
    "서열은 영구가 아니다. 중첩 0 실측 순위상관이 3개월 0.796 → 12개월 0.631"
    " → 21개월 0.570 으로 감쇠한다.",
    "BM(102110 TIGER 200)은 제외했다 — 지수 ETF 라 개별 종목 서열에 섞으면 안 된다.",
    "구간 데이터만으로 계산한다 (저장된 YZ_60 은 60일 롤링이라 연초 값이 전년을 문다).",
]


if __name__ == "__main__":
    sys.exit(main())
